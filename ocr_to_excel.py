import pymupdf
import pytesseract
import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from field_parser import parse_invoice
from validator import validate_invoice


# ---------------------------------
# Tesseract
# ---------------------------------

pytesseract.pytesseract.tesseract_cmd = (
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
)


# ---------------------------------
# 1. Open PDF
# ---------------------------------

pdf = pymupdf.open("samples/test_invoice.pdf")

page = pdf[0]

pix = page.get_pixmap(dpi=300)

image = pix.pil_image()


# ---------------------------------
# 2. OCR
# ---------------------------------

text = pytesseract.image_to_string(image)


# ---------------------------------
# 3. Parse invoice
# ---------------------------------

invoice = parse_invoice(text)


# ---------------------------------
# 4. Validate invoice
# ---------------------------------

validation = validate_invoice(invoice)

invoice["status"] = validation["status"]
invoice["confidence"] = validation["confidence"]

invoice["errors"] = "; ".join(validation["errors"])
invoice["warnings"] = "; ".join(validation["warnings"])


# ---------------------------------
# 5. Create DataFrame
# ---------------------------------

df = pd.DataFrame([invoice])


# ---------------------------------
# 6. Export Excel
# ---------------------------------

output_file = "processed_invoice.xlsx"

df.to_excel(output_file, index=False)


# ---------------------------------
# 7. Load workbook
# ---------------------------------

wb = load_workbook(output_file)

ws = wb.active


# ---------------------------------
# 8. Freeze header + filters
# ---------------------------------

ws.freeze_panes = "A2"

ws.auto_filter.ref = ws.dimensions


# ---------------------------------
# 9. Header formatting
# ---------------------------------

for cell in ws[1]:

    cell.font = Font(
        bold=True,
        color="FFFFFF"
    )

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    cell.fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )


ws.row_dimensions[1].height = 25


# ---------------------------------
# 10. Money formatting
# ---------------------------------

money_columns = [
    "taxable_amount",
    "igst",
    "total_amount"
]


headers = {
    cell.value: cell.column
    for cell in ws[1]
}


for column_name in money_columns:

    if column_name in headers:

        column_number = headers[column_name]

        for row in range(2, ws.max_row + 1):

            ws.cell(
                row=row,
                column=column_number
            ).number_format = '₹#,##0.00'


# ---------------------------------
# 11. Confidence formatting
# ---------------------------------

if "confidence" in headers:

    confidence_column = headers["confidence"]

    for row in range(2, ws.max_row + 1):

        ws.cell(
            row=row,
            column=confidence_column
        ).number_format = '0"%"'


# ---------------------------------
# 12. Status formatting
# ---------------------------------

if "status" in headers:

    status_column = headers["status"]

    for row in range(2, ws.max_row + 1):

        cell = ws.cell(
            row=row,
            column=status_column
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

        if cell.value == "Valid":

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="C6EFCE"
            )

        elif cell.value == "Warning":

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="FFEB9C"
            )

        elif cell.value == "Invalid":

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="FFC7CE"
            )


# ---------------------------------
# 13. Align cells
# ---------------------------------

for row in ws.iter_rows():

    for cell in row:

        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True
        )


# ---------------------------------
# 14. Column widths
# ---------------------------------

for column in ws.columns:

    column_letter = get_column_letter(
        column[0].column
    )

    header = str(column[0].value)

    if header in ["errors", "warnings"]:

        width = 55

    elif header in [
        "invoice_number",
        "invoice_date",
        "seller_gstin",
        "buyer_gstin"
    ]:

        width = 24

    else:

        max_length = 0

        for cell in column:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        width = max(max_length + 3, 15)

    ws.column_dimensions[
        column_letter
    ].width = width


# ---------------------------------
# 15. Save
# ---------------------------------

wb.save(output_file)


# ---------------------------------
# DONE
# ---------------------------------

print()
print("===================================")
print("OCR → PARSER → VALIDATION → EXCEL")
print("===================================")
print(f"Invoice: {invoice.get('invoice_number')}")
print(f"Status: {invoice.get('status')}")
print(f"Confidence: {invoice.get('confidence')}%")
print(f"Excel created: {output_file}")
print("===================================")