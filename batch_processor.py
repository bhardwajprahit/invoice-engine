import os



import pymupdf

import pytesseract

import pandas as pd



from pathlib import Path



from openpyxl import load_workbook

from openpyxl.styles import (

    Font,

    PatternFill,

    Alignment,

    Border,

    Side

)

from openpyxl.utils import get_column_letter



from ai_extractor import extract_invoice

from normalizer import normalize_invoice

from validator import validate_invoice

from duplicate_detector import (

    find_duplicate_files,

    calculate_file_hash

)



from database import (

    initialize_database,

    save_invoice

)





 # ============================================================

 # CONFIGURATION

 # ============================================================



TESSERACT_PATH = os.getenv(

    "TESSERACT_PATH",

    r"C:\Program Files\Tesseract-OCR\tesseract.exe"

)



SAMPLES_FOLDER = Path("samples")



OUTPUT_FILE = "batch_processed_invoices.xlsx"



SUPPORTED_FILES = {

    ".pdf",

    ".png",

    ".jpg",

    ".jpeg"

}





 # ============================================================

 # TESSERACT

 # ============================================================



pytesseract.pytesseract.tesseract_cmd = (

    TESSERACT_PATH

)





 # ============================================================

 # OCR

 # ============================================================



def ocr_pdf(file_path):



    pdf = pymupdf.open(

        file_path

    )



    pages_text = []



    try:



        for page in pdf:



            pix = page.get_pixmap(

                dpi=300

            )



            image = pix.pil_image()



            text = pytesseract.image_to_string(

                image

            )



            pages_text.append(

                text

            )



    finally:



        pdf.close()



    return "\n".join(

        pages_text

    )





def ocr_image(file_path):



    return pytesseract.image_to_string(

        str(file_path)

    )





def extract_text(file_path):



    if file_path.suffix.lower() == ".pdf":



        return ocr_pdf(

            file_path

        )



    return ocr_image(

        file_path

    )





 # ============================================================

 # PROCESS ONE DOCUMENT

 # ============================================================



def process_document(file_path):



    print(

        f"  OCR: {file_path.name}"

    )



    text = extract_text(

        file_path

    )



    print(

        "  AI extraction..."

    )



    invoice = extract_invoice(

        text

    )



    print(

        "  Normalizing..."

    )



    invoice = normalize_invoice(

        invoice

    )



    print(

        "  Validating..."

    )



    validation = validate_invoice(

        invoice

    )



    invoice["file_name"] = (

        file_path.name

    )



    invoice["status"] = (

        validation["status"]

    )



    invoice["confidence"] = (

        validation["confidence"]

    )



    invoice["review_required"] = (

        validation["review_required"]

    )



    invoice["completeness"] = (

        validation[

            "data_quality"

        ][

            "completeness"

        ]

    )



    invoice["errors"] = "; ".join(

        validation["errors"]

    )



    invoice["warnings"] = "; ".join(

        validation["warnings"]

    )



    return invoice





 # ============================================================

 # FIND DOCUMENTS

 # ============================================================



def find_documents(folder):



    if not folder.exists():



        return []



    documents = []



    for file_path in folder.iterdir():



        if not file_path.is_file():

            continue



        if file_path.suffix.lower() not in SUPPORTED_FILES:

            continue



        documents.append(

            file_path

        )



    return sorted(

        documents

    )





 # ============================================================

 # BATCH PROCESSING

 # ============================================================



def process_batch(files):



    invoices = []



    processing_errors = []



    total = len(

        files

    )



    for index, file_path in enumerate(

        files,

        start=1

    ):



        print(

            f"Processing {index}/{total}: "

            f"{file_path.name}"

        )



        try:



            invoice = process_document(

                file_path

            )



            invoices.append(

                invoice

            )



            print(

                "  Completed."

            )



        except Exception as error:



            processing_errors.append(

                {

                    "file_name": file_path.name,

                    "error": str(error)

                }

            )



            print(

                f"  ERROR: {error}"

            )



    return (

        invoices,

        processing_errors

    )





 # ============================================================

 # EXCEL COLUMNS

 # ============================================================



INVOICE_COLUMNS = [
    "invoice_number", "invoice_date", "portal_invoice_number", "portal_invoice_date",
    "currency", "seller_name", "seller_address", "seller_gstin", "buyer_name", "buyer_address",
    "buyer_gstin", "subtotal", "taxable_amount", "tax_rate", "tax_amount", "cgst", "sgst",
    "igst", "total_amount", "file_name", "status", "confidence", "review_required",
    "completeness", "errors", "warnings"
]

LINE_ITEM_COLUMNS = [
    "invoice_number", "invoice_date", "seller_name", "seller_gstin", "buyer_name",
    "product_name", "description", "model", "sku", "hsn_sac", "uom", "quantity",
    "free_quantity", "unit_price", "gross_amount", "discount", "discount_percent",
    "taxable_amount", "gst_rate", "cgst", "sgst", "item_igst", "cess", "amount", "currency",
    "file_name", "status"
]


def create_invoice_dataframe(

    invoices

):



    rows = []



    for invoice in invoices:



        row = {}



        for column in INVOICE_COLUMNS:



            row[column] = invoice.get(

                column

            )



        rows.append(

            row

        )



    return pd.DataFrame(

        rows,

        columns=INVOICE_COLUMNS

    )





 # ============================================================

 # LINE ITEM DATAFRAME

 # ============================================================



def create_line_item_dataframe(invoices):
    rows = []
    for invoice in invoices:
        items = invoice.get("line_items") or []
        if not isinstance(items, list):
            continue
        for item in items:
            if not isinstance(item, dict):
                continue
            row = {
                "invoice_number": invoice.get("invoice_number"),
                "invoice_date": invoice.get("invoice_date"),
                "seller_name": invoice.get("seller_name") or invoice.get("vendor_name"),
                "seller_gstin": invoice.get("seller_gstin"),
                "buyer_name": invoice.get("buyer_name") or invoice.get("customer_name"),
                "currency": invoice.get("currency"),
                "file_name": invoice.get("file_name"),
                "status": invoice.get("status"),
            }
            for col in LINE_ITEM_COLUMNS:
                if col not in row:
                    row[col] = item.get(col)
            row["product_name"] = row.get("product_name") or item.get("item_name") or item.get("name")
            row["model"] = row.get("model") or item.get("model_number")
            row["hsn_sac"] = row.get("hsn_sac") or item.get("hsn") or item.get("sac")
            row["item_igst"] = row.get("item_igst") if row.get("item_igst") is not None else item.get("igst")
            row["amount"] = row.get("amount") if row.get("amount") is not None else item.get("total_amount")
            rows.append(row)
    return pd.DataFrame(rows, columns=LINE_ITEM_COLUMNS)


def create_summary_dataframe(

    invoices,

    processing_errors

):



    total = len(

        invoices

    )



    valid = sum(

        invoice.get(

            "status"

        ) == "Valid"

        for invoice in invoices

    )



    warnings = sum(

        invoice.get(

            "status"

        ) == "Warning"

        for invoice in invoices

    )



    invalid = sum(

        invoice.get(

            "status"

        ) == "Invalid"

        for invoice in invoices

    )



    review_required = sum(

        invoice.get(

            "review_required",

            False

        )

        for invoice in invoices

    )



    processing_error_count = len(

        processing_errors

    )



    confidences = [



        invoice.get(

            "confidence"

        )



        for invoice in invoices



        if isinstance(

            invoice.get(

                "confidence"

            ),

            (int, float)

        )

    ]



    completeness_values = [



        invoice.get(

            "completeness"

        )



        for invoice in invoices



        if isinstance(

            invoice.get(

                "completeness"

            ),

            (int, float)

        )

    ]



    average_confidence = (



        round(

            sum(confidences)

            / len(confidences),

            1

        )



        if confidences



        else 0

    )



    average_completeness = (



        round(

            sum(completeness_values)

            / len(

                completeness_values

            ),

            1

        )



        if completeness_values



        else 0

    )



    rows = [



        {

            "Metric":

                "Documents processed",



            "Value":

                total

        },



        {

            "Metric":

                "Valid",



            "Value":

                valid

        },



        {

            "Metric":

                "Warnings",



            "Value":

                warnings

        },



        {

            "Metric":

                "Invalid",



            "Value":

                invalid

        },



        {

            "Metric":

                "Review required",



            "Value":

                review_required

        },



        {

            "Metric":

                "Processing errors",



            "Value":

                processing_error_count

        },



        {

            "Metric":

                "Average confidence",



            "Value":

                average_confidence

        },



        {

            "Metric":

                "Average completeness",



            "Value":

                average_completeness

        }

    ]



    return pd.DataFrame(

        rows

    )





 # ============================================================

 # COMMON WORKSHEET STYLE

 # ============================================================



def style_worksheet(

    worksheet,

    title

):



    title_fill = PatternFill(

        fill_type="solid",

        fgColor="1F4E78"

    )



    header_fill = PatternFill(

        fill_type="solid",

        fgColor="5B9BD5"

    )



    title_font = Font(

        bold=True,

        color="FFFFFF",

        size=15

    )



    header_font = Font(

        bold=True,

        color="FFFFFF",

        size=10

    )



    border = Border(

        bottom=Side(

            style="thin",

            color="D9E2F3"

        )

    )



    worksheet.insert_rows(

        1

    )



    worksheet["A1"] = title



    worksheet["A1"].fill = (

        title_fill

    )



    worksheet["A1"].font = (

        title_font

    )



    worksheet["A1"].alignment = (

        Alignment(

            horizontal="left",

            vertical="center"

        )

    )



    worksheet.merge_cells(

        start_row=1,

        start_column=1,

        end_row=1,

        end_column=worksheet.max_column

    )



    worksheet.row_dimensions[

        1

    ].height = 28



    for cell in worksheet[2]:



        cell.fill = (

            header_fill

        )



        cell.font = (

            header_font

        )



        cell.border = (

            border

        )



        cell.alignment = (

            Alignment(

                horizontal="center",

                vertical="center",

                wrap_text=True

            )

        )



    worksheet.row_dimensions[

        2

    ].height = 32



    for row in worksheet.iter_rows(

        min_row=3

    ):



        for cell in row:



            cell.alignment = (

                Alignment(

                    vertical="top",

                    wrap_text=True

                )

            )



    worksheet.freeze_panes = (

        "A3"

    )



    worksheet.auto_filter.ref = (

        worksheet.dimensions

    )





 # ============================================================

 # STYLE INVOICES

 # ============================================================



def style_invoices_sheet(

    worksheet

):



    style_worksheet(

        worksheet,

        "INVOICE ENGINE — INVOICES"

    )



    width_map = {



        "invoice_number": 20,

        "invoice_date": 16,

        "portal_invoice_number": 25,

        "portal_invoice_date": 20,

        "currency": 12,

        "seller_gstin": 20,

        "buyer_gstin": 20,

        "subtotal": 16,

        "taxable_amount": 18,

        "tax_rate": 12,

        "tax_amount": 16,

        "igst": 16,

        "total_amount": 18,

        "file_name": 25,

        "status": 14,

        "confidence": 14,

        "review_required": 18,

        "completeness": 15,

        "errors": 35,

        "warnings": 45

    }



    headers = [



        cell.value



        for cell in worksheet[2]

    ]



    for index, header in enumerate(

        headers,

        start=1

    ):



        worksheet.column_dimensions[

            get_column_letter(index)

        ].width = (

            width_map.get(

                header,

                18

            )

        )



    amount_columns = {



        "subtotal",

        "taxable_amount",

        "tax_amount",

        "igst",

        "total_amount"

    }



    for row in worksheet.iter_rows(

        min_row=3

    ):



        for cell in row:



            header = (

                worksheet.cell(

                    row=2,

                    column=cell.column

                ).value

            )



            if header in amount_columns:



                cell.number_format = (

                    "#,##0.00"

                )



            elif header in {



                "confidence",

                "completeness"



            }:



                cell.number_format = (

                    "0"

                )





 # ============================================================

 # STYLE LINE ITEMS

 # ============================================================



def style_line_items_sheet(

    worksheet

):



    style_worksheet(

        worksheet,

        "INVOICE ENGINE — LINE ITEMS"

    )



    width_map = {



        "invoice_number": 20,

        "invoice_date": 16,

        "description": 50,

        "quantity": 12,

        "unit_price": 16,

        "amount": 16,

        "currency": 12,

        "file_name": 25

    }



    headers = [



        cell.value



        for cell in worksheet[2]

    ]



    for index, header in enumerate(

        headers,

        start=1

    ):



        worksheet.column_dimensions[

            get_column_letter(index)

        ].width = (

            width_map.get(

                header,

                18

            )

        )



    for row in worksheet.iter_rows(

        min_row=3

    ):



        for cell in row:



            header = (

                worksheet.cell(

                    row=2,

                    column=cell.column

                ).value

            )



            if header in {



                "unit_price",

                "amount"



            }:



                cell.number_format = (

                    "#,##0.00"

                )





 # ============================================================

 # STYLE SUMMARY

 # ============================================================



def style_summary_sheet(

    worksheet

):



    title_fill = PatternFill(

        fill_type="solid",

        fgColor="1F4E78"

    )



    header_fill = PatternFill(

        fill_type="solid",

        fgColor="5B9BD5"

    )



    title_font = Font(

        bold=True,

        color="FFFFFF",

        size=18

    )



    header_font = Font(

        bold=True,

        color="FFFFFF",

        size=11

    )



    label_font = Font(

        bold=True,

        size=11

    )



    worksheet.insert_rows(

        1

    )



    worksheet["A1"] = (

        "INVOICE ENGINE — PROCESSING SUMMARY"

    )



    worksheet["A1"].fill = (

        title_fill

    )



    worksheet["A1"].font = (

        title_font

    )



    worksheet["A1"].alignment = (

        Alignment(

            horizontal="left",

            vertical="center"

        )

    )



    worksheet.merge_cells(

        "A1:B1"

    )



    worksheet.row_dimensions[

        1

    ].height = 32



    for cell in worksheet[2]:



        cell.fill = (

            header_fill

        )



        cell.font = (

            header_font

        )



        cell.alignment = (

            Alignment(

                horizontal="center"

            )

        )



    for row in worksheet.iter_rows(

        min_row=3

    ):



        worksheet.cell(

            row=row[0].row,

            column=1

        ).font = (

            label_font

        )



        for cell in row:



            cell.alignment = (

                Alignment(

                    vertical="center"

                )

            )



    worksheet.column_dimensions[

        "A"

    ].width = 45



    worksheet.column_dimensions[

        "B"

    ].width = 25



    worksheet.freeze_panes = (

        "A3"

    )





 # ============================================================

 # CREATE EXCEL

 # ============================================================



def create_excel(invoices, processing_errors):
    """Create a formatted workbook with four useful sheets."""
    if not invoices and not processing_errors:
        return False
    invoice_df = create_invoice_dataframe(invoices)
    line_item_df = create_line_item_dataframe(invoices)
    summary_df = create_summary_dataframe(invoices, processing_errors)
    issues_df = pd.DataFrame(
        [{"file_name": e.get("file_name", ""), "error": e.get("error", str(e))}
         for e in (processing_errors or [])], columns=["file_name", "error"]
    )
    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        summary_df.to_excel(writer, index=False, sheet_name="Summary")
        invoice_df.to_excel(writer, index=False, sheet_name="Invoices")
        line_item_df.to_excel(writer, index=False, sheet_name="Line Items")
        issues_df.to_excel(writer, index=False, sheet_name="Processing Issues")
    workbook = load_workbook(OUTPUT_FILE)
    for ws in workbook.worksheets:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A3" if ws.title == "Summary" else "A2"
        if ws.max_row >= 2:
            ws.auto_filter.ref = ws.dimensions
        header_row = 2 if ws.title == "Summary" else 1
        for cell in ws[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="183153")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if ws.title == "Summary":
            for cell in ws[1]:
                cell.font = Font(bold=True, size=16, color="183153")
            ws.row_dimensions[1].height = 30
        ws.row_dimensions[header_row].height = 30
        for idx in range(1, ws.max_column + 1):
            letter = get_column_letter(idx)
            width = 12
            for cells in ws.iter_rows(min_col=idx, max_col=idx):
                val = cells[0].value
                if val is not None:
                    width = max(width, min(len(str(val)) + 2, 34))
            ws.column_dimensions[letter].width = width
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (1, 2))
    workbook.save(OUTPUT_FILE)
    return True


def print_summary(

    invoices,

    processing_errors,

    excel_created,

    duplicate_files

):



    valid_count = sum(

        invoice.get(

            "status"

        ) == "Valid"

        for invoice in invoices

    )



    warning_count = sum(

        invoice.get(

            "status"

        ) == "Warning"

        for invoice in invoices

    )



    invalid_count = sum(

        invoice.get(

            "status"

        ) == "Invalid"

        for invoice in invoices

    )



    print()



    print(

        "==================================="

    )



    print(

        "INVOICE ENGINE"

    )



    print(

        "BATCH PROCESSING COMPLETE"

    )



    print(

        "==================================="

    )



    print(

        f"Documents processed: "

        f"{len(invoices)}"

    )



    print(

        f"Duplicates skipped: "

        f"{len(duplicate_files)}"

    )



    print(

        f"Valid: "

        f"{valid_count}"

    )



    print(

        f"Warnings: "

        f"{warning_count}"

    )



    print(

        f"Invalid: "

        f"{invalid_count}"

    )



    print(

        f"Processing errors: "

        f"{len(processing_errors)}"

    )



    if excel_created:



        print(

            f"Excel created: "

            f"{OUTPUT_FILE}"

        )



    else:



        print(

            "Excel created: No"

        )



    print(

        "==================================="

    )





 # ============================================================

 # MAIN

 # ============================================================



def main():



    initialize_database()



    files = find_documents(

        SAMPLES_FOLDER

    )



    print(

        f"Found {len(files)} "

        f"supported documents."

    )



    if not files:



        print(

            "No supported documents found."

        )



        return



    (

        unique_files,

        duplicate_files

    ) = find_duplicate_files(

        files

    )



    print(

        f"Unique documents: "

        f"{len(unique_files)}"

    )



    print(

        f"Duplicate documents: "

        f"{len(duplicate_files)}"

    )



    for duplicate in duplicate_files:



        print(

            f"Duplicate: "

            f"{duplicate['file_name']} "

            f"-> "

            f"{duplicate['duplicate_of']}"

        )



    (

        invoices,

        processing_errors

    ) = process_batch(

        unique_files

    )



    for invoice in invoices:



        file_name = invoice.get(

            "file_name"

        )



        matching_file = next(

            (

                file_path



                for file_path in unique_files



                if file_path.name == file_name

            ),

            None

        )



        if matching_file is None:

            continue



        try:



            file_hash = (

                calculate_file_hash(

                    matching_file

                )

            )



            save_invoice(

                invoice,

                file_hash

            )



        except Exception as error:



            processing_errors.append(

                {

                    "file_name": file_name,

                    "error":

                        f"Database error: {error}"

                }

            )



            print(

                f"  DATABASE ERROR: "

                f"{error}"

            )



    excel_created = create_excel(

        invoices,

        processing_errors

    )



    print_summary(

        invoices,

        processing_errors,

        excel_created,

        duplicate_files

    )





if __name__ == "__main__":



    main()