
from pathlib import Path
import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment


# Find all invoice files
files = list(Path("samples").glob("*.txt"))

all_invoices = []


# Map invoice labels to dictionary keys
field_map = {
    "Invoice Number:": "invoice_number",
    "Invoice Date:": "invoice_date",
    "Vendor:": "vendor",
    "GSTIN:": "gstin",
    "Item:": "item",
    "Quantity:": "quantity",
    "Unit Price:": "unit_price",
    "Subtotal:": "subtotal",
    "GST:": "gst",
    "Total Amount:": "total_amount",
}


# Process every invoice
for file_path in files:
    text = file_path.read_text()
    lines = text.splitlines()

    invoice = {}

    # Extract fields
    for line in lines:
        for label, key in field_map.items():
            if line.startswith(label):
                value = line.split(":", 1)[1].strip()
                invoice[key] = value

    # Convert numeric values
    invoice["quantity"] = int(invoice["quantity"])
    invoice["unit_price"] = float(invoice["unit_price"])
    invoice["subtotal"] = float(invoice["subtotal"])
    invoice["gst"] = float(invoice["gst"])
    invoice["total_amount"] = float(invoice["total_amount"])

    # Validate invoice calculations
    subtotal_ok = (
        invoice["quantity"] * invoice["unit_price"]
        == invoice["subtotal"]
    )

    total_ok = (
        invoice["subtotal"] + invoice["gst"]
        == invoice["total_amount"]
    )

    if subtotal_ok and total_ok:
        invoice["status"] = "Valid"
    elif not subtotal_ok:
        invoice["status"] = "Subtotal mismatch"
    else:
        invoice["status"] = "Total mismatch"

    all_invoices.append(invoice)


# Create table
df = pd.DataFrame(all_invoices)


# Create Excel file
df.to_excel("processed_invoices.xlsx", index=False)


# Open Excel file for formatting
wb = load_workbook("processed_invoices.xlsx")
ws = wb.active


# Freeze header row
ws.freeze_panes = "A2"


# Enable filters
ws.auto_filter.ref = ws.dimensions


# Make headers bold and centered
for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center")


# Format money columns
for row in ws.iter_rows(min_row=2):
    for cell in row:
        if cell.column in [7, 8, 9, 10]:
            cell.number_format = '₹#,##0.00'


# Set column widths
for column in ws.columns:
    column_letter = column[0].column_letter

    if column_letter in ["G", "H", "I", "J"]:
        ws.column_dimensions[column_letter].width = 18
    else:
        max_length = 0

        for cell in column:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[column_letter].width = max_length + 3


# Save final Excel file
wb.save("processed_invoices.xlsx")


print("Invoice processing complete.")
print(f"Processed {len(all_invoices)} invoices.")
print("Excel file created: processed_invoices.xlsx")